#!/usr/bin/env python3
"""Build the week's posting plan: day folders with the files in posting order, a POST-TODAY.md per day,
and campaigns/AD-SCHEDULE.xlsx (one row per post, thumbnail, clickable link to the file, caption, status).

Usage: python3 tools/build_week.py
Edit WEEK below to change the plan. Files listed with "src" are copied into campaigns/<date>/ as "dest".
"""
import shutil
import subprocess
from pathlib import Path

import imageio_ffmpeg
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
C = ROOT / "campaigns"
NICO = "https://www.amazon.com/dp/B0HLVTTK7T"
BENNY = "https://www.amazon.com/dp/B0HLVSLXK7"
LINK = {"Nico": NICO, "Benny": BENNY}


def tt(text, tags):
    return f"{text} Link in bio\n{tags}"


def meta(text, book, tags):
    return f"{text}\nGet it on Amazon: {LINK[book]} (Instagram: link in bio)\n{tags}"


def yt(title, desc, book, tags):
    return f"TITLE:\n{title}\n\nDESCRIPTION:\n{desc}\nGet it on Amazon: {LINK[book]}\n{tags}"


def pin(title, desc, book, board):
    return f"TITLE: {title}\nDESCRIPTION: {desc}\nDESTINATION LINK: {LINK[book]}\nBOARD: {board}"


STORY_TAGS = "#bedtimestory #christmascoloring #kidsbooks #storytime #toddleractivities"
B_TAGS = "#coloringbook #colorwithme #satisfying #toddleractivities #screenfree"
N_TAGS = "#christmascoloring #coloringbook #colorwithme #toddleractivities #stockingstuffers"
NV = "campaigns/nico/videos/"
BV = "campaigns/benny/videos/"

# (time, platforms, book, format, src (repo path, file or folder), dest name, caption, note)
WEEK = {
    ("2026-10-07", "Wed"): [
        ("19:00", "IG + FB", "Nico", "Reel: bedtime story PART 3 (the end)", NV + "bedtime-story-part3-voice-music.mp4",
         "1-nico-story-part3.mp4",
         meta("Bedtime story: Nico's Snowy Day, the end 🌙🦌 After the best snowman ever, Mama tucks Nico into his warm bed. "
              "Goodnight, Nico! The whole story is a Christmas coloring book kids color AND read, ages 3-6.", "Nico",
              STORY_TAGS), "Story with link after posting"),
        ("19:00", "TikTok", "Nico", "Video: bedtime story part 2", NV + "bedtime-story-part2-voice-music.mp4",
         "2-nico-story-part2.mp4",
         tt("Christmas bedtime story for kids: Nico's Snowy Day, part 2 ⛄ What did Nico and Tilly build? "
            "Part 3 tomorrow, follow so you don't miss it!", STORY_TAGS), "Pin comment: Did you guess it was a snowman? ⛄👇"),
        ("19:00", "YouTube Shorts", "Nico", "Short: bedtime story part 2", NV + "bedtime-story-part2-voice-music.mp4",
         "2-nico-story-part2.mp4",
         yt("Christmas bedtime story for kids: Nico's Snowy Day, part 2 ⛄ #shorts #storytime",
            "Nico and Tilly build something big in the snow! Part 3 tomorrow. The whole story is a color & read Christmas "
            "book for kids ages 3-6.", "Nico", "#bedtimestory #christmascoloring #kidsbooks"), "Not made for kids"),
        ("21:00", "TikTok", "Benny", "Video: guess the colors (kite)", BV + "guess-kite-voice.mp4", "3-benny-guess-kite.mp4",
         tt("Coloring page for kids: guess the colors before the end 🪁🧸 Benny and Bella fly a kite. Kids color every page "
            "AND read the sentence. Ages 3-6, best with crayons.", B_TAGS), "Pin comment: What color did you guess for the kite? 👇"),
        ("21:00", "IG + FB", "Benny", "Reel: guess the colors (kite)", BV + "guess-kite-voice.mp4", "3-benny-guess-kite.mp4",
         meta("Guess the colors before the end! 🪁🧸 Benny and Bella fly a kite high up in the sky. Kids color every page AND "
              "read the short sentence. Ages 3-6.", "Benny", "#coloringbook #kidsactivities #screenfree #toddleractivities"), ""),
        ("21:00", "YouTube Shorts", "Benny", "Short: guess the colors (kite)", BV + "yt-guess-kite-voice.mp4",
         "3-yt-benny-guess-kite.mp4",
         yt("Guess the colors before the end! 🪁 #shorts #coloringbook",
            "Benny the Bear's Cozy Day: a color & read story book for kids ages 3-6.", "Benny",
            "#coloringbook #kidsactivities #satisfying"), "Not made for kids"),
        ("any", "Pinterest", "Benny", "Image pin: before/after", "campaigns/2026-10-05/pinterest/3-benny-pin-before-after.png",
         "4-pin-benny-before-after.png",
         pin("Cute Bear Coloring Pages for Kids 3-6: Big Easy Pages",
             "Before and after: a cute bear coloring page. Benny the Bear's Cozy Day is a color and read story book for ages "
             "3-6. Quiet time activity before bed.", "Benny", "Cozy Animal Stories"), ""),
    ],
    ("2026-10-08", "Thu"): [
        ("19:00", "IG + FB", "Nico", "Reel: 30 minutes. Zero screens.", NV + "screenfree-voice-music.mp4",
         "1-nico-screenfree.mp4",
         meta("30 minutes. Zero screens. 🎄 If you have a 3-6 year old, try this instead of the tablet: a Christmas story "
              "they color AND read. Send this to a parent who needs a quiet hour 💛", "Nico",
              "#screenfree #toddleractivities #christmascoloring #coloringbook #kidsactivities"), "Story with link"),
        ("19:00", "TikTok", "Nico", "Video: bedtime story part 3 (the end)", NV + "bedtime-story-part3-voice-music.mp4",
         "2-nico-story-part3.mp4",
         tt("Christmas bedtime story for kids: Nico's Snowy Day, the end 🌙 Goodnight, Nico! The whole story is a coloring "
            "book kids color AND read.", STORY_TAGS), ""),
        ("19:00", "YouTube Shorts", "Nico", "Short: bedtime story part 3 (the end)", NV + "bedtime-story-part3-voice-music.mp4",
         "2-nico-story-part3.mp4",
         yt("Christmas bedtime story for kids: Nico's Snowy Day, the end 🌙 #shorts #storytime",
            "Goodnight, Nico! The whole story is a color & read Christmas book for kids ages 3-6.", "Nico",
            "#bedtimestory #christmascoloring #kidsbooks"), "Not made for kids"),
        ("20:00", "IG + FB + TikTok", "Benny", "Photo carousel (6 slides): gift", "campaigns/benny/carousel-gift",
         "3-benny-gift-carousel",
         meta("Birthday gift idea for kids 3-6 🎁 A story they color AND read: Benny the Bear's Cozy Day. Big pictures, big "
              "letters. Swipe to peek inside 👉", "Benny", "#giftsforkids #birthdaygift #coloringbook #kidsactivities"),
         "Post slides 1 to 6 in order. TikTok: photo post + calm sound, caption ends 'Link in bio'"),
        ("21:00", "TikTok", "Benny", "Video: can your 4-year-old read this?", BV + "can-they-read-voice.mp4",
         "4-benny-can-they-read.mp4",
         tt("Can your 4-year-old read this page? 📖 One big picture + one short sentence on every page, so kids color it AND "
            "read it.", "#learningtoread #preschoolathome #toddleractivities #coloringbook #bedtimestory"), ""),
        ("21:00", "IG + FB", "Benny", "Reel: can your 4-year-old read this?", BV + "can-they-read-voice.mp4",
         "4-benny-can-they-read.mp4",
         meta("Can your 4-year-old read this page? 📖 Every page: one big picture to color + one short sentence in big "
              "letters. Color it, then read it together.", "Benny", "#learningtoread #preschoolathome #toddleractivities"), ""),
        ("21:00", "YouTube Shorts", "Benny", "Short: guess the colors (hide and seek)", BV + "yt-guess-hide-voice.mp4",
         "5-yt-benny-guess-hide.mp4",
         yt("Guess the colors: hide and seek with Benny 🐻 #shorts #coloringbook",
            "Benny the Bear's Cozy Day: a color & read story book for kids ages 3-6.", "Benny",
            "#coloringbook #kidsactivities #satisfying"), "Not made for kids"),
        ("any", "Pinterest", "Nico", "Image pin: inside pages", "campaigns/nico/images/pin2-pages.png", "6-pin-nico-pages.png",
         pin("Easy Christmas Coloring Pages for Kids: Color + Read",
             "Inside Nico the Little Reindeer's Snowy Day: big, easy winter coloring pages with one short sentence each. "
             "For kids ages 3-6.", "Nico", "Christmas Gifts for Kids"), ""),
    ],
    ("2026-10-09", "Fri"): [
        ("19:00", "TikTok", "Nico", "Video: guess the colors (snowman scarf)", NV + "guess-scarf-voice.mp4",
         "1-nico-guess-scarf.mp4",
         tt("Christmas coloring book for kids: guess the colors before the end ⛄🖍️ Nico gives the snowman his red scarf.",
            N_TAGS), "Pin comment: What color did you guess for the scarf? 👇"),
        ("19:00", "IG + FB", "Nico", "Reel: guess the colors (snowman scarf)", NV + "guess-scarf-voice.mp4",
         "1-nico-guess-scarf.mp4",
         meta("Guess the colors before the end! ⛄🖍️ Nico gives the snowman his red scarf. Kids color every page AND read "
              "the short sentence. A stocking stuffer that isn't candy, ages 3-6.", "Nico",
              "#christmascoloring #coloringbook #stockingstuffers"), "Story with link"),
        ("19:00", "YouTube Shorts", "Nico", "Short: guess the colors (snowman scarf)", NV + "yt-guess-scarf-voice.mp4",
         "1-yt-nico-guess-scarf.mp4",
         yt("Christmas coloring page: guess the colors! ⛄ #shorts #coloringbook",
            "Nico the Little Reindeer's Snowy Day: a color & read Christmas story book for kids ages 3-6.", "Nico",
            "#christmascoloring #coloringbook #stockingstuffers"), "Not made for kids"),
        ("21:00", "TikTok", "Nico", "Video: 30 minutes. Zero screens.", NV + "screenfree-voice-music.mp4",
         "2-nico-screenfree.mp4",
         tt("30 minutes. Zero screens. 🎄 If you have a 3-6 year old, try this instead of the tablet: a Christmas story they "
            "color AND read.", "#screenfree #toddleractivities #christmascoloring #coloringbook #momlife"), ""),
        ("21:00", "YouTube Shorts", "Nico", "Short: 30 minutes. Zero screens.", NV + "screenfree-voice-music.mp4",
         "2-nico-screenfree.mp4",
         yt("30 minutes. Zero screens. Try this instead of the tablet 🎄 #shorts #screenfree",
            "A Christmas story kids ages 3-6 color AND read: Nico the Little Reindeer's Snowy Day.", "Nico",
            "#screenfree #christmascoloring #kidsactivities"), "Not made for kids"),
        ("21:00", "IG + FB", "Benny", "Reel: guess the colors (hide and seek)", BV + "guess-hide-voice.mp4",
         "3-benny-guess-hide.mp4",
         meta("Guess the colors before the end! 🐻🐰 Benny and Bella play hide and seek. Kids color every page AND read the "
              "story. Ages 3-6.", "Benny", "#coloringbook #kidsactivities #screenfree #colorwithme"), ""),
        ("any", "Pinterest", "Nico", "Image pin: gift", "campaigns/nico/images/pin3-gift.png", "4-pin-nico-gift.png",
         pin("Stocking Stuffer for Toddlers: Reindeer Coloring Story Book",
             "A Christmas gift for kids 3-6 that isn't more plastic: a color and read story book with a cute reindeer.",
             "Nico", "Stocking Stuffers"), ""),
    ],
    ("2026-10-10", "Sat"): [
        ("15:00", "TikTok + IG + FB + YouTube", "Nico", "Carousel video: gift (voice)", NV + "carousel-gift-voice.mp4",
         "1-nico-gift-carousel-video.mp4",
         meta("Stocking stuffer that isn't candy? 🎁🦌 A Christmas coloring book kids color AND read. 29 winter pages, ages "
              "3-6.", "Nico", "#stockingstuffers #christmasgifts #coloringbook #giftsforkids"),
         "TikTok: end the caption with 'Link in bio'. YouTube title: Stocking stuffer that isn't candy 🎁 #shorts"),
        ("15:30", "IG + FB + TikTok", "Nico", "Photo carousel (6 slides): bedtime story", "campaigns/nico/carousel-story",
         "2-nico-story-carousel",
         meta("A bedtime story kids color AND read 🌙🦌 Swipe to peek inside 👉 One big picture and one short sentence on every "
              "page. Ages 3-6.", "Nico", "#bedtimestory #christmascoloring #coloringbook #kidsbooks"),
         "Slides 1 to 6 in order. TikTok: photo post + calm sound"),
        ("16:00", "TikTok", "Benny", "Video: guess the colors (butterfly)", BV + "guess-butterfly-voice.mp4",
         "3-benny-guess-butterfly.mp4",
         tt("Coloring page for kids: guess the colors before the end 🦋🧸 A butterfly lands on Benny's nose.", B_TAGS),
         "Pin comment: What color did you guess for the butterfly? 👇"),
        ("16:00", "IG + FB", "Benny", "Reel: guess the colors (butterfly)", BV + "guess-butterfly-voice.mp4",
         "3-benny-guess-butterfly.mp4",
         meta("Guess the colors before the end! 🦋🧸 A butterfly lands on Benny's nose. Hello! Kids color every page AND read "
              "the story. Ages 3-6.", "Benny", "#coloringbook #kidsactivities #screenfree"), ""),
        ("16:00", "YouTube Shorts", "Benny", "Short: guess the colors (butterfly)", BV + "yt-guess-butterfly-voice.mp4",
         "3-yt-benny-guess-butterfly.mp4",
         yt("Guess the colors: Benny and the butterfly 🦋 #shorts #coloringbook",
            "Benny the Bear's Cozy Day: a color & read story book for kids ages 3-6.", "Benny",
            "#coloringbook #kidsactivities #satisfying"), "Not made for kids"),
        ("any", "Pinterest", "Benny", "Image pin: inside pages", "campaigns/benny/images/pin2-pages.png",
         "4-pin-benny-pages.png",
         pin("Cute Bear Coloring Pages for Kids 3-6",
             "Inside Benny the Bear's Cozy Day: big, easy coloring pages with one short sentence each. Quiet time activity.",
             "Benny", "Screen-Free Fun"), ""),
    ],
    ("2026-10-11", "Sun"): [
        ("15:00", "TikTok", "Nico", "Video: guess the colors (snow angel)", NV + "guess-angel-voice.mp4",
         "1-nico-guess-angel.mp4",
         tt("Christmas coloring book for kids: guess the colors before the end ❄️🦌 Nico makes his very first snow angel.",
            N_TAGS), "Pin comment: What color did you guess for Nico? 👇"),
        ("15:00", "IG + FB", "Nico", "Reel: guess the colors (snow angel)", NV + "guess-angel-voice.mp4",
         "1-nico-guess-angel.mp4",
         meta("Guess the colors before the end! ❄️🦌 Nico makes his very first snow angel. Kids color every page AND read the "
              "story. Ages 3-6.", "Nico", "#christmascoloring #coloringbook #stockingstuffers"), "Story with link"),
        ("15:00", "YouTube Shorts", "Nico", "Short: guess the colors (snow angel)", NV + "yt-guess-angel-voice.mp4",
         "1-yt-nico-guess-angel.mp4",
         yt("Christmas coloring page: guess the colors! ❄️ #shorts #coloringbook",
            "Nico the Little Reindeer's Snowy Day: a color & read Christmas story book for kids ages 3-6.", "Nico",
            "#christmascoloring #coloringbook"), "Not made for kids"),
        ("16:00", "TikTok", "Benny", "Video: read-along (kite, hill, hide and seek)", BV + "read-kite-voice.mp4",
         "2-benny-read-along.mp4",
         tt("Can your 4-year-old read this? 📖 Read along with Benny: one short sentence on every page.",
            "#learningtoread #preschoolathome #toddleractivities #coloringbook"), ""),
        ("16:00", "IG + FB", "Benny", "Reel: read-along", BV + "read-kite-voice.mp4", "2-benny-read-along.mp4",
         meta("Can your 4-year-old read this? 📖 Read along with Benny: one big picture and one short sentence on every page.",
              "Benny", "#learningtoread #preschoolathome #toddleractivities"), ""),
        ("16:00", "YouTube Shorts", "Benny", "Short: read-along", BV + "yt-read-kite-voice.mp4", "2-yt-benny-read-along.mp4",
         yt("Can your 4 year old read this? 📖 #shorts #learningtoread",
            "Read along with Benny the Bear: one big picture + one short sentence in big letters on every page. Ages 3-6.",
            "Benny", "#learningtoread #preschool #coloringbook"), "Not made for kids"),
        ("any", "Pinterest", "Nico", "Image pin: cover", "campaigns/nico/images/pin1-cover.png", "3-pin-nico-cover.png",
         pin("Cute Reindeer Christmas Coloring Book for Kids 3-6",
             "Nico the Little Reindeer's Snowy Day: a color and read Christmas story book for preschool and kindergarten.",
             "Nico", "Christmas Gifts for Kids"), ""),
        ("evening", "Sunday check", "All", "KDP sales + best 3 posts", "", "",
         "Send Claude a screenshot of KDP Reports + the 3 posts with the most views/link clicks", ""),
    ],
}

# Monday + Tuesday: already made (packs in campaigns/2026-10-05 and 2026-10-06), listed for the record
M, T = "campaigns/2026-10-05/", "campaigns/2026-10-06/"
DONE = [
    ("2026-10-05", "Mon", "10:30", "Instagram", "Nico", "Reel: guess the colors", M + "instagram/1-nico-guess-colors.mp4", "Posted"),
    ("2026-10-05", "Mon", "19:00", "YouTube Shorts", "Nico", "Short: guess the colors", M + "youtube/1-yt-nico-guess-colors.mp4", "Scheduled"),
    ("2026-10-05", "Mon", "19:00", "TikTok", "Nico", "Video: guess the colors", M + "tiktok/1-nico-guess-colors.mp4", "Scheduled"),
    ("2026-10-05", "Mon", "19:00", "Facebook", "Nico", "Reel: guess the colors", M + "facebook/1-nico-guess-colors.mp4", "Scheduled"),
    ("2026-10-05", "Mon", "20:00", "IG + FB", "Nico", "Photo carousel: gift", "campaigns/nico/carousel-gift", "Scheduled"),
    ("2026-10-05", "Mon", "20:00", "TikTok", "Nico", "Photo carousel: gift", "campaigns/nico/carousel-gift", "Scheduled"),
    ("2026-10-05", "Mon", "21:00", "TikTok", "Benny", "Video: same page 3 ways", M + "tiktok/3-benny-3-ways.mp4", "Scheduled"),
    ("2026-10-05", "Mon", "21:00", "YouTube Shorts", "Benny", "Short: can your 4-year-old read this?", M + "youtube/2-yt-benny-can-they-read.mp4", "Scheduled"),
    ("2026-10-05", "Mon", "21:00", "IG + FB", "Nico", "Reel: bedtime story part 1", M + "instagram/4-nico-bedtime-story-part1.mp4", "Scheduled"),
    ("2026-10-06", "Tue", "19:00", "IG + FB", "Nico", "Reel: bedtime story part 2", T + "instagram-facebook/1-nico-story-part2.mp4", "Planned"),
    ("2026-10-06", "Tue", "19:00", "TikTok", "Nico", "Video: bedtime story part 1", T + "tiktok/1-nico-story-part1.mp4", "Planned"),
    ("2026-10-06", "Tue", "19:00", "YouTube Shorts", "Nico", "Short: bedtime story part 1", T + "youtube/1-nico-story-part1.mp4", "Planned"),
    ("2026-10-06", "Tue", "21:00", "IG + FB", "Benny", "Reel: can you find Bella?", T + "instagram-facebook/2-benny-find-bella.mp4", "Planned"),
    ("2026-10-06", "Tue", "21:00", "TikTok", "Benny", "Video: can you find Bella?", T + "tiktok/2-benny-find-bella.mp4", "Planned"),
    ("2026-10-06", "Tue", "21:00", "YouTube Shorts", "Benny", "Short: can you find Bella?", T + "youtube/2-yt-benny-find-bella.mp4", "Planned"),
    ("2026-10-06", "Tue", "any", "Pinterest", "Nico", "Image pin: before/after", T + "pinterest/1-nico-pin-before-after.png", "Planned"),
]


def thumb(path, out):
    """Small preview picture: a frame from 40% into a video, slide 1 of a carousel, or the image itself."""
    p = ROOT / path
    if p.is_dir():
        p = p / "slide1.png"
    if p.suffix == ".mp4":
        ff = imageio_ffmpeg.get_ffmpeg_exe()
        info = subprocess.run([ff, "-i", str(p)], capture_output=True, text=True).stderr
        import re
        h, m, s = re.search(r"Duration: (\d+):(\d+):([\d.]+)", info).groups()
        t = (int(h) * 3600 + int(m) * 60 + float(s)) * 0.4
        subprocess.run([ff, "-loglevel", "error", "-y", "-ss", f"{t:.2f}", "-i", str(p), "-frames:v", "1",
                        "-vf", "scale=96:-1", str(out)], check=True)
    else:
        im = Image.open(p).convert("RGB")
        im.thumbnail((96, 170))
        im.save(out)
    return out


def main():
    rows = []
    for (date, day), posts in WEEK.items():
        folder = C / date
        folder.mkdir(exist_ok=True)
        lines = [f"# {day} {date}: what to post (open AD-SCHEDULE.xlsx for the whole week)\n",
                 "Videos already have voice + soft music: **don't add another sound**. "
                 "Instagram + Facebook: Meta Business Suite, tick both.\n"]
        for time, plat, book, fmt, src, dest, cap, note in posts:
            rel = ""
            if src:
                s, d = ROOT / src, folder / dest
                if s.is_dir():
                    shutil.copytree(s, d, dirs_exist_ok=True)
                elif not d.exists() or d.stat().st_size != s.stat().st_size:
                    shutil.copy2(s, d)
                rel = f"campaigns/{date}/{dest}"
            rows.append((date, day, time, plat, book, fmt, rel, cap, note, "Planned"))
            lines.append(f"\n## {time} · {plat} · {dest or fmt}\n{fmt}" + (f" · {note}" if note else "") +
                         (f"\n```\n{cap}\n```\n" if cap else "\n"))
        (folder / "POST-TODAY.md").write_text("".join(lines))
    done = [(d, dy, t, p, b, f, src, "In that day's CAPTIONS.md / POST-TODAY.md", "", st)
            for d, dy, t, p, b, f, src, st in DONE]
    write_xlsx(done + rows)


def write_xlsx(rows):
    tmp = Path("/tmp/claude-thumbs")
    tmp.mkdir(exist_ok=True)
    F = lambda **k: Font(name="Arial", **k)  # noqa: E731
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws["A1"] = "Little Crayon Tales: ad schedule, week of Oct 5"
    ws["A1"].font = F(bold=True, size=14)
    ws["A2"] = ("Click 'Open file' to open the video (open this sheet from the kids-coloring-books/campaigns folder after "
                "'pull'). Yellow = you fill in. Captions are ready to copy.")
    ws["A2"].font = F(italic=True, size=9)
    hdr = ["Date", "Day", "Time (CZ)", "Platform", "Book", "Preview", "What", "Open file", "Caption / title (copy)",
           "Notes", "Status", "Views (3 days)", "Link clicks"]
    ws.append(hdr)
    head = PatternFill("solid", fgColor="2F4A6D")
    for c in ws[3]:
        c.font = F(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    shade = {"Mon": "EAF2FB", "Tue": "FFFFFF", "Wed": "EAF2FB", "Thu": "FFFFFF", "Fri": "EAF2FB", "Sat": "FFFFFF",
             "Sun": "EAF2FB"}
    yellow = PatternFill("solid", fgColor="FFF4C2")
    line = Border(bottom=Side(style="thin", color="CCCCCC"))
    for i, (date, day, time, plat, book, fmt, rel, cap, note, status) in enumerate(rows):
        r = 4 + i
        vals = [date, day, time, plat, book, "", fmt, "", cap, note, status, "", ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = F(size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = PatternFill("solid", fgColor=shade[day])
            cell.border = line
        for c in (11, 12, 13):
            ws.cell(r, c).fill = yellow
        if rel:
            link = ws.cell(r, 8, "▶ Open file" if not (ROOT / rel).is_dir() else "📁 Open slides folder")
            link.hyperlink = str(Path(rel).relative_to("campaigns"))
            link.font = F(size=10, color="0645AD", underline="single")
            ws.cell(r, 10).value = (note + "  |  " if note else "") + rel
            t = thumb(rel, tmp / f"t{r}.png")
            img = XLImage(str(t))
            ws.add_image(img, f"F{r}")
            ws.row_dimensions[r].height = 135
        else:
            ws.row_dimensions[r].height = 60
    for col, w in zip("ABCDEFGHIJKLM", [11, 5, 8, 16, 7, 15, 26, 14, 62, 30, 11, 10, 10]):
        ws.column_dimensions[col].width = w
    n = 3 + len(rows)
    dv = DataValidation(type="list", formula1='"Planned,To do,Scheduled,Posted,Skipped"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K4:K{n}")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:M{n}"

    t = wb.create_sheet("Totals")
    t["A1"] = "Posts per platform (updates when you change Status)"
    t["A1"].font = F(bold=True, size=13)
    t.append(["Platform", "Planned", "Scheduled", "Posted", "Total"])
    for c in t[2]:
        c.font = F(bold=True, color="FFFFFF")
        c.fill = head
    plats = [("TikTok", ["TikTok"]), ("YouTube Shorts", ["YouTube"]), ("Instagram", ["Instagram", "IG"]),
             ("Facebook", ["Facebook", "FB"]), ("Pinterest", ["Pinterest"])]
    for i, (p, names) in enumerate(plats, 3):
        t.cell(i, 1, p).font = F(size=11)
        for j, st in enumerate(["Planned", "Scheduled", "Posted"], 2):
            f = "+".join(f'COUNTIFS(Schedule!$D$4:$D${n},"*{x}*",Schedule!$K$4:$K${n},"{st}")' for x in names)
            t.cell(i, j, "=" + f).font = F(size=11)
        t.cell(i, 5, f"=SUM(B{i}:D{i})").font = F(size=11)
    t["A9"] = "A row like 'TikTok + IG + FB + YouTube' counts once for each platform it names."
    t["A9"].font = F(italic=True, size=9)
    t.column_dimensions["A"].width = 20

    w = wb.create_sheet("Workplan")
    w["A1"] = "Daily workplan (from the ad research, see campaigns/WORKPLAN.md)"
    w["A1"].font = F(bold=True, size=13)
    w.append(["Platform", "Per day", "Carousels", "Extra"])
    for c in w[2]:
        c.font = F(bold=True, color="FFFFFF")
        c.fill = head
    for row in [["TikTok", "2 videos (19:00 + 21:00)", "+1 photo carousel Mon/Thu/Sat", "Pin a question comment"],
                ["YouTube Shorts", "2 Shorts (yt- files)", "Carousel only as a video", "'Not made for kids'"],
                ["Instagram", "2 Reels, one per book (19:00 + 21:00)", "+1 carousel post Mon/Thu/Sat", "Story with Link sticker"],
                ["Facebook", "2 Reels (same as Instagram)", "+1 carousel post Mon/Thu/Sat", "Business Suite = IG + FB at once"],
                ["Pinterest", "1 pin", "-", "Destination link = Amazon"],
                ["Times (Czech)", "Mon-Fri 19:00-21:00", "Sat-Sun 15:00-16:00", ""],
                ["Which book", "Nico 2 of every 3 posts until Dec 10", "New books: 1 video a day for 7 days", "Dex: on hold"]]:
        w.append(row)
    for row in w.iter_rows(min_row=3):
        for c in row:
            c.font = F(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, wd in zip("ABCD", [16, 34, 34, 34]):
        w.column_dimensions[col].width = wd
    wb.save(C / "AD-SCHEDULE.xlsx")
    print("✓", C / "AD-SCHEDULE.xlsx", len(rows), "rows")


if __name__ == "__main__":
    main()
