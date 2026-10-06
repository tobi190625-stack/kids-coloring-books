#!/usr/bin/env python3
"""One easy posting sheet per day: every post in time order, with a picture of the video, a click-to-open
file link, the platform, the caption to copy, the comment to pin, and a Done column.

Usage: python tools/day_sheet.py campaigns/2026-10-07     (reads posts.json there, writes POSTS.xlsx there)
posts.json: {"day": "...", "posts": [{"time", "platform", "book", "file", "caption", optional "title",
"pin", "link", "note", "status"}]}
"""
import json
import subprocess
import sys
import tempfile
from pathlib import Path

import imageio_ffmpeg
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image

COLORS = {"TikTok": "FFD6E4", "YouTube Shorts": "FFD9D9", "Instagram + Facebook": "E3DBFF", "Pinterest": "FFE4D1"}
COLS = [("Done", 8), ("Time", 8), ("Platform", 16), ("Book", 8), ("Picture", 16), ("File (click to open)", 24),
        ("Title (YouTube / Pinterest)", 30), ("Caption: copy all of it", 70), ("Pin this comment", 24),
        ("Notes", 30)]


def thumb(src, out):
    """Small picture of a video (frame at 2.5 s) or an image, 110 px wide."""
    if src.suffix.lower() == ".mp4":
        subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-loglevel", "error", "-y", "-ss", "2.5", "-i", str(src),
                        "-frames:v", "1", "-update", "1", str(out)], check=True)
        im = Image.open(out)
    else:
        im = Image.open(src)
    im = im.convert("RGB")
    im.thumbnail((110, 196))
    im.save(out)
    return im.size


def main(day_dir):
    day = Path(day_dir)
    data = json.loads((day / "posts.json").read_text(encoding="utf-8"))
    posts = sorted(data["posts"], key=lambda p: (p["time"] == "any", p["time"], p["platform"]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Posts"
    ws["A1"] = f"{data['day']}: what to post"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A2"] = ("Click the file name to open it. Copy the caption into the app. Type ✓ in Done when posted. "
                "Videos already have sound: don't add music.")
    ws["A2"].font = Font(italic=True, color="555555")
    head = 4
    thin = Side(style="thin", color="BBBBBB")
    for c, (name, w) in enumerate(COLS, 1):
        cell = ws.cell(head, c, name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3B4A63")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = ws.cell(head + 1, 1)
    done = DataValidation(type="list", formula1='"✓"', allow_blank=True)
    ws.add_data_validation(done)
    tmp = Path(tempfile.mkdtemp())
    for i, p in enumerate(posts):
        r = head + 1 + i
        status = p.get("status", "")
        vals = ["✓" if status.startswith("SCHEDULED") else "", p["time"], p["platform"], p["book"], "", p["file"],
                p.get("title", ""), p["caption"] + (f"\n\nDestination link: {p['link']}" if p.get("link") else ""),
                p.get("pin", ""), " ".join(x for x in (status.replace("SCHEDULED", "Already scheduled by Claude"),
                                                          p.get("note", "")) if x)]
        fill = PatternFill("solid", fgColor=COLORS.get(p["platform"], "FFFFFF"))
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if c in (2, 3, 4):
                cell.fill = fill
                cell.font = Font(bold=True)
        done.add(ws.cell(r, 1))
        ws.cell(r, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(r, 1).font = Font(size=18, bold=True, color="2E9E4F")
        f = ws.cell(r, 6)
        f.hyperlink = p["file"]
        f.font = Font(color="0563C1", underline="single")
        src = day / p["file"]
        if src.exists():
            t = tmp / f"t{i}.png"
            w, h = thumb(src, t)
            img = XLImage(str(t))
            ws.add_image(img, f"E{r}")
            ws.row_dimensions[r].height = max(h * 0.75 + 8, 120)
        else:
            f.value = f"{p['file']} (MISSING)"
    out = day / "POSTS.xlsx"
    wb.save(out)
    print("✓", out, len(posts), "posts")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
